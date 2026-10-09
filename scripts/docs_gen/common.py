# -*- coding: utf-8 -*-
"""AutoLink Pro — Kit de style commun pour la génération des documents Word."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime, os

TODAY = datetime.date(2026, 10, 9).strftime('%d/%m/%Y')
DOCS_DIR = os.path.join(os.path.dirname(__file__), '..', '..', 'documentation')
ASSETS   = os.path.join(DOCS_DIR, 'assets')

# ── Palette AutoLink Pro ─────────────────────────────────────────────────────
C_TEAL  = RGBColor(0x0D, 0x94, 0x88)   # primary-600 (marque)
C_SKY   = RGBColor(0x0E, 0xA5, 0xE9)
C_ORA   = RGBColor(0xF9, 0x73, 0x16)
C_GRN   = RGBColor(0x10, 0xB9, 0x81)
C_RED   = RGBColor(0xEF, 0x44, 0x44)
C_NAVY  = RGBColor(0x0F, 0x17, 0x2A)
C_DARK  = RGBColor(0x1E, 0x29, 0x3B)
C_GREY  = RGBColor(0x64, 0x74, 0x8B)
C_WHT   = RGBColor(0xFF, 0xFF, 0xFF)

H_TEAL = '0D9488'; H_NAVY = '0F172A'; H_SKY = '0EA5E9'; H_LIGHT = 'ECFEFF'
H_GREYBG = 'F1F5F9'; H_ORA = 'F97316'

# ── helpers bas niveau ───────────────────────────────────────────────────────
def _shd(cell, hex6):
    pr = cell._tc.get_or_add_tcPr()
    e = OxmlElement('w:shd')
    e.set(qn('w:val'), 'clear'); e.set(qn('w:color'), 'auto'); e.set(qn('w:fill'), hex6)
    pr.append(e)

def _set_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), 'CBD5E1')
        borders.append(el)
    tblPr.append(borders)

def new_doc(landscape=False):
    doc = Document()
    s = doc.sections[0]
    if landscape:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21)
    else:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2.2)
    # Pied de page : numéro de page + nom projet
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run('AutoLink Pro — ')
    r.font.size = Pt(8); r.font.color.rgb = C_GREY
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
    r2 = footer.add_run('  ·  Diffusion restreinte')
    r2.font.size = Pt(8); r2.font.color.rgb = C_GREY
    return doc

def cover(doc, num, doc_type, subtitle, note=None):
    """Page de garde stylée : bandeau couleur, titre, sous-titre, statut."""
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    r = p.add_run('AUTOLINK PRO'); r.bold = True; r.font.size = Pt(16)
    r.font.color.rgb = C_GREY; r.font.name = 'Arial'
    doc.add_paragraph()
    t = doc.add_table(rows=1, cols=1); t.alignment = 1
    c = t.cell(0, 0); _shd(c, H_NAVY)
    cp = c.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run(f'\n{doc_type}\n'); cr.bold = True; cr.font.size = Pt(26); cr.font.color.rgb = C_WHT
    doc.add_paragraph()
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle); r2.font.size = Pt(14); r2.font.color.rgb = C_TEAL; r2.italic = True
    doc.add_paragraph(); doc.add_paragraph()
    for line in [f'Référence : {num}   ·   Version 1.0   ·   {TODAY}',
                 'Plateforme : https://autolink-pro.worldwide-international.business',
                 '"Votre mobilité, notre mission" — Douala & Yaoundé, Cameroun']:
        p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(line); r3.font.size = Pt(10.5); r3.font.color.rgb = C_GREY
    if note:
        doc.add_paragraph()
        callout(doc, note, fill='FFF7ED')
    doc.add_page_break()

def callout(doc, text, fill=H_LIGHT, bold_prefix=None):
    t = doc.add_table(rows=1, cols=1)
    _set_borders(t)
    c = t.cell(0, 0); _shd(c, fill)
    p = c.paragraphs[0]
    if bold_prefix:
        r = p.add_run(bold_prefix + '  '); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = C_NAVY
    r = p.add_run(text); r.font.size = Pt(10); r.font.color.rgb = C_DARK
    doc.add_paragraph()

def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '14')
    bot.set(qn('w:space'), '3'); bot.set(qn('w:color'), H_TEAL)
    b.append(bot); pPr.append(b)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = C_NAVY

def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    r = p.add_run('▸ ' + text); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = C_TEAL

def h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = C_DARK

def para(doc, text, bold=False, italic=False, size=10.5, color=C_DARK, center=False):
    p = doc.add_paragraph()
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); r.bold = bold; r.italic = italic
    r.font.size = Pt(size); r.font.color.rgb = color
    return p

def bullets(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style='List Bullet')
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = C_NAVY
            r2 = p.add_run(it[1]); r2.font.size = Pt(size); r2.font.color.rgb = C_DARK
        else:
            r = p.add_run(it); r.font.size = Pt(size); r.font.color.rgb = C_DARK

def numbered(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style='List Number')
        r = p.add_run(it); r.font.size = Pt(size); r.font.color.rgb = C_DARK

def table(doc, headers, rows, widths=None, header_fill=H_NAVY, size=9.5):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    _set_borders(t)
    for j, h in enumerate(headers):
        c = t.cell(0, j); _shd(c, header_fill)
        r = c.paragraphs[0].add_run(h)
        r.bold = True; r.font.size = Pt(size); r.font.color.rgb = C_WHT
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i + 1, j)
            if i % 2 == 1: _shd(c, H_GREYBG)
            p = c.paragraphs[0]
            r = p.add_run(str(v)); r.font.size = Pt(size); r.font.color.rgb = C_DARK
    if widths:
        for j, w in enumerate(widths):
            for i in range(len(rows) + 1):
                t.cell(i, j).width = Cm(w)
    doc.add_paragraph()
    return t

def img(doc, filename, width_cm=15.5, caption=None):
    path = os.path.join(ASSETS, filename)
    if not os.path.exists(path):
        para(doc, f'[Schéma à générer : {filename}]', italic=True, color=C_GREY, center=True)
        return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm))
    if caption:
        para(doc, caption, italic=True, size=9, color=C_GREY, center=True)

def status_block(doc, pct, remaining):
    """Bloc « marge de réalisation » : avancement + actions pour 100 %."""
    h2(doc, 'Marge de réalisation du document')
    table(doc, ['Avancement', 'Reste à faire pour 100 %'],
          [[f'{pct} %', remaining]], widths=[3.5, 13])
